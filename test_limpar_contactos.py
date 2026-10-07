"""
Testes. Correr com:
    python -m unittest -v
"""

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from limpar_contactos import (limpar_contactos, limpar_data, limpar_email,
                              limpar_nome, limpar_telefone, proteger_texto)


class TestFuncoes(unittest.TestCase):
    def test_nome(self):
        self.assertEqual(limpar_nome("  MARIA   DA silva "), ("Maria da Silva", ""))
        self.assertEqual(limpar_nome("   ")[1], "nome em branco")

    def test_email(self):
        self.assertEqual(limpar_email("  ANA@Exemplo.PT "), ("ana@exemplo.pt", ""))
        self.assertEqual(limpar_email("ana@exemplo")[1], "email inválido")
        self.assertEqual(limpar_email("ana silva@exemplo.pt")[1], "email inválido")

    def test_telefone(self):
        esperado = "+351 900 123 456"
        for entrada in ["900123456", "900 123 456", "+351900123456",
                        "00351 900123456", "(+351) 900-123-456"]:
            self.assertEqual(limpar_telefone(entrada), (esperado, ""), entrada)
        self.assertEqual(limpar_telefone("12345")[1], "telefone não parece português")
        self.assertEqual(limpar_telefone(""), ("", ""))

    def test_data(self):
        for entrada in ["05/03/2025", "2025-03-05", "5-3-2025", "05.03.2025"]:
            self.assertEqual(limpar_data(entrada), ("05/03/2025", ""), entrada)
        self.assertEqual(limpar_data("32/13/2025")[1], "data em formato desconhecido")
        self.assertEqual(limpar_data("01/01/2999")[1], "data no futuro")

    def test_proteger_texto(self):
        self.assertEqual(proteger_texto("=1+1"), "'=1+1")
        self.assertEqual(proteger_texto("+351 900 123 456"), "+351 900 123 456")
        self.assertEqual(proteger_texto("+351900123456=HACK"), "'+351900123456=HACK")


class TestFicheiro(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def criar(self, texto: str) -> Path:
        caminho = self.pasta / "c.csv"
        caminho.write_text(texto, encoding="utf-8")
        return caminho

    def test_duplicados_e_rejeitados(self):
        entrada = self.criar(
            "nome;email;telefone;data_registo;localidade\n"
            "Ana Silva;ana@exemplo.pt;900123456;01/01/2025;Ovar\n"
            "ANA SILVA; ANA@EXEMPLO.PT ;;;Ovar\n"
            ";sem@exemplo.pt;;;\n"
            "Rui Costa;rui@exemplo;;;\n"
        )
        r = limpar_contactos(entrada, self.pasta / "s.xlsx")
        self.assertEqual(r["lidas"], 4)
        self.assertEqual(r["limpos"], 1)
        self.assertEqual(r["rejeitados"], 3)

    def test_original_nao_e_alterado(self):
        texto = "nome;email;telefone;data_registo;localidade\nana;ana@exemplo.pt;;;\n"
        entrada = self.criar(texto)
        limpar_contactos(entrada, self.pasta / "s.xlsx")
        self.assertEqual(entrada.read_text(encoding="utf-8"), texto)

    def test_saida_igual_a_entrada_e_recusada(self):
        entrada = self.pasta / "c.xlsx"
        pd.DataFrame({"nome": ["a"], "email": ["a@b.pt"], "telefone": [""],
                      "data_registo": [""], "localidade": [""]}).to_excel(entrada, index=False)
        with self.assertRaises(ValueError):
            limpar_contactos(entrada, entrada)

    def test_formula_fica_como_texto(self):
        entrada = self.criar("nome;email;telefone;data_registo;localidade\n"
                             "=1+1;f@exemplo.pt;;;@SUM(A1)\n")
        saida = self.pasta / "s.xlsx"
        limpar_contactos(entrada, saida)
        valores = [c.value for linha in load_workbook(saida)["Contactos"].iter_rows() for c in linha]
        self.assertFalse(any(isinstance(v, str) and v.startswith(("=", "@")) for v in valores))

    def test_extensao_recusada(self):
        caminho = self.pasta / "macro.xlsm"
        caminho.write_bytes(b"x")
        with self.assertRaises(ValueError):
            limpar_contactos(caminho, self.pasta / "s.xlsx")


if __name__ == "__main__":
    unittest.main()
